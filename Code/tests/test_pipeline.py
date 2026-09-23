import datetime as dt
import json
import logging
from pathlib import Path

import pytest
from pptx import Presentation

import main as cli
from src.console import LOGGER_NAMES
from src.pipeline import ReportResult, generate_report, resolve_output_path
from src.rendering.slides.notes import NOTES_SHAPE_NAME

TODAY = dt.date(2026, 9, 22)
ROWS = [
    {"Date": "2026-09-20", "CampaignName": "Поиск_Бренд", "Impressions": 4000, "Clicks": 120, "Cost": 5000.0, "Conversions": 5},
    {"Date": "2026-09-20", "CampaignName": "РСЯ_Интересы", "Impressions": 15000, "Clicks": 300, "Cost": 4000.0, "Conversions": 6},
    {"Date": "2026-09-21", "CampaignName": "Поиск_Бренд", "Impressions": 4200, "Clicks": 130, "Cost": 5200.0, "Conversions": 4},
    {"Date": "2026-09-21", "CampaignName": "РСЯ_Интересы", "Impressions": 16000, "Clicks": 310, "Cost": 4100.0, "Conversions": 7},
]
NOTES_MD = """# Итоги месяца

Расход вырос на **8 %** при стабильном CPA.

- Отключили площадки РСЯ с высоким CPA
- Добавили минус-слова

1. Запустить A/B-тест объявлений
2. Снизить ставки на брендовые запросы
"""


@pytest.fixture(autouse=True)
def reset_console_logging():
    """CLI подключает обработчик логов к stderr теста — снимаем его, чтобы не писать в закрытый поток."""
    yield
    for name in LOGGER_NAMES:
        logger = logging.getLogger(name)
        for handler in [h for h in logger.handlers if h.get_name() == "report-console"]:
            logger.removeHandler(handler)
        logger.setLevel(logging.NOTSET)


@pytest.fixture
def csv_file(tmp_path: Path) -> Path:
    path = tmp_path / "export.csv"
    lines = ["Date,CampaignName,Impressions,Clicks,Cost,Conversions"]
    lines += [",".join(str(value) for value in row.values()) for row in ROWS]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


@pytest.fixture
def json_file(tmp_path: Path) -> Path:
    path = tmp_path / "export.json"
    path.write_text(json.dumps(ROWS, ensure_ascii=False), encoding="utf-8")
    return path


def notes_texts(path: Path) -> list[str]:
    deck = Presentation(path)
    return [shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.name == NOTES_SHAPE_NAME]


def notes_xml(path: Path) -> str:
    deck = Presentation(path)
    shapes = [shape for slide in deck.slides for shape in slide.shapes if shape.name == NOTES_SHAPE_NAME]
    return "".join(shape._element.xml for shape in shapes)


# --- Пайплайн ---------------------------------------------------------------


def test_mock_run_builds_report_with_timings(tmp_path):
    result = generate_report(mock=True, output=tmp_path, mock_seed=1, today=TODAY)

    assert isinstance(result, ReportResult)
    assert result.output_path == tmp_path.resolve() / "report_ООО_Доставка_Плюс_2026-08-23_2026-09-21.pptx"
    assert result.output_path.exists()
    assert (result.records_count, result.campaigns_count, result.slides_count) == (4 * 30, 4, 5)
    assert result.period == (dt.date(2026, 8, 23), dt.date(2026, 9, 21))
    assert [step.name for step in result.timings] == ["Конфиг", "Данные", "Агрегация", "Рендеринг"]
    assert result.total_seconds > 0
    assert result.timing_summary().split(" · ")[-1].startswith("всего")
    assert not result.has_notes
    assert notes_texts(result.output_path) == ["Вставьте комментарии специалиста перед отправкой клиенту..."]


@pytest.mark.parametrize("source", ["csv_file", "json_file"])
def test_file_inputs(tmp_path, request, source):
    result = generate_report(input_path=request.getfixturevalue(source), output=tmp_path / "r.pptx", today=TODAY)

    assert result.output_path == (tmp_path / "r.pptx").resolve()
    assert (result.records_count, result.campaigns_count) == (4, 2)
    assert result.period == (dt.date(2026, 9, 20), dt.date(2026, 9, 21))


def test_markdown_notes_become_headings_paragraphs_and_lists(tmp_path):
    result = generate_report(mock=True, notes=NOTES_MD, output=tmp_path / "r.pptx", mock_seed=1, today=TODAY)

    [text] = notes_texts(result.output_path)
    assert result.has_notes
    assert text.splitlines() == [
        "Итоги месяца",
        "Расход вырос на 8 % при стабильном CPA.",
        "Отключили площадки РСЯ с высоким CPA",
        "Добавили минус-слова",
        "Запустить A/B-тест объявлений",
        "Снизить ставки на брендовые запросы",
    ]
    xml = notes_xml(result.output_path)
    assert xml.count("<a:buChar") == 2 and xml.count("<a:buAutoNum") == 2


def test_notes_from_file_and_from_items(tmp_path):
    notes_file = tmp_path / "notes.md"
    notes_file.write_text(NOTES_MD, encoding="utf-8")

    from_file = generate_report(mock=True, notes=notes_file, output=tmp_path / "a.pptx", today=TODAY)
    from_items = generate_report(mock=True, notes=["Первый пункт", "  ", "Второй пункт"], output=tmp_path / "b.pptx")

    assert notes_texts(from_file.output_path)[0].startswith("Итоги месяца")
    assert notes_texts(from_items.output_path) == ["Первый пункт\nВторой пункт"]


def test_notes_file_in_cp1251_is_read_with_warning(tmp_path, caplog):
    notes_file = tmp_path / "notes.txt"
    notes_file.write_bytes("Выводы в старой кодировке".encode("cp1251"))

    with caplog.at_level(logging.WARNING, logger="src"):
        result = generate_report(mock=True, notes=notes_file, output=tmp_path, today=TODAY)

    assert notes_texts(result.output_path) == ["Выводы в старой кодировке"]
    assert "cp1251" in caplog.text


def test_blank_notes_keep_placeholder_and_warn(tmp_path, caplog):
    with caplog.at_level(logging.WARNING, logger="src"):
        result = generate_report(mock=True, notes="   \n  ", output=tmp_path, today=TODAY)

    assert not result.has_notes
    assert "пустые" in caplog.text


def test_long_notes_continue_on_next_slides(tmp_path):
    items = [f"Пункт {index}: скорректировали ставки и минус-слова, CPA снизился на {index} %" for index in range(40)]

    result = generate_report(mock=True, notes=items, output=tmp_path, today=TODAY)

    deck = Presentation(result.output_path)
    titles = [shape.text_frame.text for slide in deck.slides for shape in slide.shapes if shape.name == "Заголовок"]
    assert result.slides_count == len(deck.slides) > 5
    assert titles[-1] == "Выводы и план оптимизации (продолжение)"
    assert sum(text.count("Пункт ") for text in notes_texts(result.output_path)) == 40  # ничего не потерялось


def test_warns_about_incomplete_today_and_period_mismatch(tmp_path, csv_file, caplog):
    with caplog.at_level(logging.WARNING, logger="src"):
        generate_report(input_path=csv_file, output=tmp_path, today=dt.date(2026, 9, 21))

    assert "неполная" in caplog.text
    assert "period_days: 30" in caplog.text


def test_missing_input_file(tmp_path):
    with pytest.raises(FileNotFoundError, match="Файл данных не найден"):
        generate_report(input_path=tmp_path / "nope.csv", output=tmp_path)


def test_missing_config_and_notes_file(tmp_path):
    with pytest.raises(FileNotFoundError, match="Конфиг не найден"):
        generate_report(tmp_path / "nope.yaml", mock=True, output=tmp_path)
    with pytest.raises(FileNotFoundError, match="Файл заметок не найден"):
        generate_report(mock=True, notes=tmp_path / "nope.md", output=tmp_path)


def test_invalid_config_names_the_file(tmp_path):
    config = tmp_path / "bad.yaml"
    config.write_text("report_metadata:\n  client_name: Клиент\n  clinet: опечатка\nslides: []\n", encoding="utf-8")

    with pytest.raises(ValueError, match="bad.yaml не прошёл валидацию"):
        generate_report(config, mock=True, output=tmp_path)


@pytest.mark.parametrize("kwargs", [{}, {"mock": True, "input_path": "export.csv"}])
def test_requires_exactly_one_data_source(tmp_path, kwargs):
    with pytest.raises(ValueError, match="источник данных"):
        generate_report(output=tmp_path, **kwargs)


def test_resolve_output_path(tmp_path):
    period = (dt.date(2026, 8, 23), dt.date(2026, 9, 21))

    assert resolve_output_path(tmp_path / "r.pptx", "Клиент", period) == tmp_path / "r.pptx"
    assert resolve_output_path(tmp_path, "ООО «Доставка Плюс»", period).name == (
        "report_ООО_Доставка_Плюс_2026-08-23_2026-09-21.pptx"
    )
    with pytest.raises(ValueError, match=".pptx"):
        resolve_output_path(tmp_path / "report.pdf", "Клиент", period)


# --- CLI --------------------------------------------------------------------


def report_path_from(stdout: str) -> Path:
    first_line = stdout.splitlines()[0]
    return Path(first_line.split(": ", 1)[1])


def test_cli_mock_run(tmp_path, capsys):
    code = cli.main(["--mock", "--seed", "1", "-o", str(tmp_path), "--no-color"])

    out, err = capsys.readouterr()
    assert code == 0
    assert out.startswith("Отчёт готов:")
    assert report_path_from(out).exists()
    assert "[4/4] Рендеринг" in err and "INFO" in err
    assert "\033[" not in out + err  # --no-color


def test_cli_csv_with_inline_notes(tmp_path, csv_file, capsys):
    target = tmp_path / "report.pptx"

    code = cli.main(["-i", str(csv_file), "-n", r"- Первый вывод\n- Второй вывод", "-o", str(target), "--no-color"])

    assert code == 0
    assert notes_texts(target) == ["Первый вывод\nВторой вывод"]
    assert "выводы специалиста добавлены" in capsys.readouterr().out


def test_cli_json_with_notes_file(tmp_path, json_file, capsys):
    notes_file = tmp_path / "notes.md"
    notes_file.write_text(NOTES_MD, encoding="utf-8")

    code = cli.main(["-i", str(json_file), "--notes-file", str(notes_file), "-o", str(tmp_path), "--no-color"])

    assert code == 0
    assert notes_texts(report_path_from(capsys.readouterr().out))[0].startswith("Итоги месяца")


def test_cli_missing_input_file(tmp_path, capsys):
    code = cli.main(["-i", str(tmp_path / "nope.csv"), "-o", str(tmp_path), "--no-color"])

    out, err = capsys.readouterr()
    assert code == 1
    assert "ERROR" in err and "Файл данных не найден" in err
    assert out == ""


def test_cli_requires_data_source(capsys):
    with pytest.raises(SystemExit) as exit_info:
        cli.main([])

    assert exit_info.value.code == 2
    assert "укажите источник данных" in capsys.readouterr().err


def test_cli_rejects_both_sources(capsys):
    with pytest.raises(SystemExit) as exit_info:
        cli.main(["-i", "export.csv", "--mock"])

    assert exit_info.value.code == 2


def test_cli_warns_that_seed_needs_mock(tmp_path, csv_file, capsys):
    code = cli.main(["-i", str(csv_file), "--seed", "3", "-o", str(tmp_path), "--no-color"])

    assert code == 0
    assert "только вместе с --mock" in capsys.readouterr().err
