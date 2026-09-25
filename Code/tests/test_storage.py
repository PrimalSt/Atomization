import datetime as dt
import logging
import sqlite3
from pathlib import Path

import pytest
from pptx import Presentation

import main as cli
from src.console import LOGGER_NAMES
from src.pipeline import generate_report
from src.rendering.formatting import NBSP
from src.schemas import CampaignSummary
from src.storage import (
    SCHEMA_VERSION,
    get_monthly_metrics,
    get_previous_month_metrics,
    init_storage,
    is_full_month,
    month_key,
    parse_month,
    previous_month,
    report_month_for,
    save_monthly_metrics,
    snapshot_metrics,
)

CLIENT = "ООО Доставка Плюс"


@pytest.fixture(autouse=True)
def reset_console_logging():
    """CLI подключает обработчик логов к stderr теста — снимаем его после теста."""
    yield
    for name in LOGGER_NAMES:
        logger = logging.getLogger(name)
        for handler in [h for h in logger.handlers if h.get_name() == "report-console"]:
            logger.removeHandler(handler)
        logger.setLevel(logging.NOTSET)


@pytest.fixture
def db(tmp_path: Path) -> Path:
    return tmp_path / "local_storage" / "history.db"


def summary(campaign_id: int = 1, *, cost: float, clicks: int, impressions: int, conversions: int) -> CampaignSummary:
    return CampaignSummary(
        campaign_id=campaign_id,
        campaign_name=f"Кампания {campaign_id}",
        date_from=dt.date(2026, 8, 1),
        date_to=dt.date(2026, 8, 31),
        days_count=31,
        impressions=impressions,
        clicks=clicks,
        cost=cost,
        conversions=conversions,
    )


AUGUST = [
    summary(1, cost=30_000.0, clicks=600, impressions=20_000, conversions=20),
    summary(2, cost=10_000.0, clicks=400, impressions=30_000, conversions=5),
]


# --- Месяцы ---------------------------------------------------------------------


def test_month_helpers():
    assert month_key(dt.date(2026, 9, 21)) == "2026-09"
    assert parse_month("2026-09") == (2026, 9)
    assert previous_month("2026-09") == "2026-08"
    assert previous_month("2026-01") == "2025-12"


@pytest.mark.parametrize("month", ["2026-13", "2026-9", "09.2026", "", "2026-00"])
def test_invalid_month_is_rejected(month):
    with pytest.raises(ValueError, match="YYYY-MM"):
        parse_month(month)


@pytest.mark.parametrize(
    ("period", "expected"),
    [
        ((dt.date(2026, 8, 1), dt.date(2026, 8, 31)), "2026-08"),
        ((dt.date(2026, 8, 23), dt.date(2026, 9, 21)), "2026-09"),  # скользящие 30 дней: больше дней в сентябре
        ((dt.date(2026, 8, 17), dt.date(2026, 9, 13)), "2026-08"),
        ((dt.date(2026, 8, 30), dt.date(2026, 9, 2)), "2026-09"),  # поровну — более поздний месяц
        ((dt.date(2025, 12, 20), dt.date(2026, 1, 15)), "2026-01"),
    ],
)
def test_report_month_is_the_month_with_most_days(period, expected):
    assert report_month_for(period) == expected


def test_is_full_month():
    assert is_full_month((dt.date(2026, 2, 1), dt.date(2026, 2, 28)))
    assert not is_full_month((dt.date(2026, 8, 2), dt.date(2026, 8, 31)))
    assert not is_full_month((dt.date(2026, 8, 1), dt.date(2026, 9, 30)))


# --- Сохранение и чтение ------------------------------------------------------------


def test_save_creates_database_and_returns_snapshot(db):
    saved = save_monthly_metrics(CLIENT, "2026-08", AUGUST, db_path=db)

    assert db.is_file()  # папка local_storage создана
    assert saved["client_id"] == CLIENT and saved["report_month"] == "2026-08"
    assert (saved["total_spend"], saved["total_clicks"], saved["total_impressions"], saved["total_conversions"]) == (
        40_000.0,
        1_000,
        50_000,
        25,
    )
    assert saved["avg_cpa"] == pytest.approx(1_600.0)  # от сумм: 40 000 / 25
    assert saved["avg_ctr"] == pytest.approx(2.0)  # 1 000 / 50 000 × 100
    assert saved["created_at"] == saved["updated_at"]
    assert dt.datetime.fromisoformat(saved["created_at"]).tzinfo is not None
    assert get_monthly_metrics(CLIENT, "2026-08", db_path=db) == saved


def test_table_schema(db):
    init_storage(db)

    with sqlite3.connect(db) as connection:
        columns = [row[1] for row in connection.execute("PRAGMA table_info(monthly_snapshots)")]
        version = connection.execute("PRAGMA user_version").fetchone()[0]
    assert columns == [
        "client_id", "report_month", "total_spend", "total_clicks", "total_impressions",
        "total_conversions", "avg_cpa", "avg_ctr", "created_at", "updated_at",
    ]  # fmt: skip
    assert version == SCHEMA_VERSION


def test_upsert_replaces_month_and_keeps_created_at(db, monkeypatch):
    monkeypatch.setattr("src.storage._now", lambda: "2026-09-01T10:00:00+03:00")
    first = save_monthly_metrics(CLIENT, "2026-08", AUGUST, db_path=db)
    monkeypatch.setattr("src.storage._now", lambda: "2026-09-02T12:30:00+03:00")
    corrected = [summary(1, cost=31_000.0, clicks=620, impressions=21_000, conversions=21)]

    second = save_monthly_metrics(CLIENT, "2026-08", corrected, db_path=db)

    assert (second["total_spend"], second["total_conversions"]) == (31_000.0, 21)
    assert second["created_at"] == first["created_at"] == "2026-09-01T10:00:00+03:00"
    assert second["updated_at"] == "2026-09-02T12:30:00+03:00"
    with sqlite3.connect(db) as connection:
        assert connection.execute("SELECT COUNT(*) FROM monthly_snapshots").fetchone()[0] == 1


def test_previous_month_metrics(db):
    save_monthly_metrics(CLIENT, "2026-08", AUGUST, db_path=db)

    previous = get_previous_month_metrics(CLIENT, "2026-09", db_path=db)

    assert previous is not None and previous["report_month"] == "2026-08"
    assert get_previous_month_metrics(CLIENT, "2026-08", db_path=db) is None  # июля нет
    assert get_previous_month_metrics(CLIENT, "2026-10", db_path=db) is None  # через пропуск не сравниваем


def test_previous_month_across_year_boundary(db):
    save_monthly_metrics(CLIENT, "2025-12", AUGUST, db_path=db)

    assert get_previous_month_metrics(CLIENT, "2026-01", db_path=db)["report_month"] == "2025-12"


def test_clients_are_isolated(db):
    save_monthly_metrics(CLIENT, "2026-08", AUGUST, db_path=db)

    assert get_previous_month_metrics("ООО Другой клиент", "2026-09", db_path=db) is None
    assert get_previous_month_metrics(f"  {CLIENT} ", "2026-09", db_path=db) is not None  # пробелы по краям не важны


def test_undefined_averages_are_stored_as_null(db):
    idle = [summary(cost=0.0, clicks=0, impressions=0, conversions=0)]

    saved = save_monthly_metrics(CLIENT, "2026-08", idle, db_path=db)

    assert saved["avg_cpa"] is None and saved["avg_ctr"] is None
    assert snapshot_metrics(saved).cpa is None


def test_snapshot_metrics_recomputes_all_kpis(db):
    saved = save_monthly_metrics(CLIENT, "2026-08", AUGUST, db_path=db)

    metrics = snapshot_metrics(saved)

    assert (metrics.cost, metrics.clicks, metrics.impressions, metrics.conversions) == (40_000.0, 1_000, 50_000, 25)
    assert metrics.cpc == pytest.approx(40.0)
    assert metrics.cr == pytest.approx(2.5)


def test_reading_missing_database_does_not_create_it(db):
    assert get_previous_month_metrics(CLIENT, "2026-09", db_path=db) is None
    assert not db.exists()


@pytest.mark.parametrize(
    ("client", "month", "rows", "message"),
    [
        ("  ", "2026-08", AUGUST, "client_id"),
        (CLIENT, "август", AUGUST, "YYYY-MM"),
        (CLIENT, "2026-08", [], "сохранять нечего"),
    ],
)
def test_save_validates_input(db, client, month, rows, message):
    with pytest.raises(ValueError, match=message):
        save_monthly_metrics(client, month, rows, db_path=db)
    assert not db.exists()


def test_database_from_newer_version_is_refused(db):
    db.parent.mkdir(parents=True)
    with sqlite3.connect(db) as connection:
        connection.execute(f"PRAGMA user_version = {SCHEMA_VERSION + 1}")

    with pytest.raises(sqlite3.DatabaseError, match="более новой версией"):
        save_monthly_metrics(CLIENT, "2026-08", AUGUST, db_path=db)


# --- Динамика MoM в отчёте ------------------------------------------------------------

HEADER = "Date,CampaignName,Impressions,Clicks,Cost,Conversions"


def month_csv(path: Path, month: int, *, cost: float, clicks: int, conversions: int) -> Path:
    days = 31 if month == 8 else 30
    lines = [HEADER]
    lines += [
        f"2026-{month:02d}-{day:02d},Поиск_Бренд,1000,{clicks},{cost},{conversions}" for day in range(1, days + 1)
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def kpi_captions(path: Path) -> dict[str, str]:
    """Пояснения под карточками KPI; неразрывные пробелы заменены обычными для читаемости проверок."""
    deck = Presentation(path)
    return {
        shape.name.split(":")[0].removeprefix("KPI "): shape.text_frame.text.replace(NBSP, " ")
        for slide in deck.slides
        for shape in slide.shapes
        if shape.name.endswith(": пояснение")
    }


def test_second_month_report_shows_month_over_month(tmp_path, db):
    august = month_csv(tmp_path / "aug.csv", 8, cost=1000.0, clicks=20, conversions=2)
    september = month_csv(tmp_path / "sep.csv", 9, cost=1100.0, clicks=25, conversions=2)
    today = dt.date(2026, 10, 5)

    first = generate_report(input_path=august, output=tmp_path / "aug.pptx", history_db=db, today=today)
    second = generate_report(input_path=september, output=tmp_path / "sep.pptx", history_db=db, today=today)

    assert (first.report_month, first.compared_to, first.history_saved) == ("2026-08", None, True)
    assert (second.report_month, second.compared_to, second.history_saved) == ("2026-09", "2026-08", True)
    assert kpi_captions(first.output_path)["avg_cpa"] == "расход ÷ конверсии"  # без прошлого месяца — как раньше
    captions = kpi_captions(second.output_path)
    # Расход: 33 000 против 31 000 → +6,5 %; конверсии: 60 против 62 → −3,2 %;
    # CPA: 550 против 500 → +10 %; CTR: 2,5 % против 2,0 % → +0,5 п.п.
    assert captions == {
        "total_spend": "▲ 6,5% к августу",
        "conversions": "▼ 3,2% к августу",
        "avg_cpa": "▲ 10,0% к августу",
        "avg_ctr": "▲ 0,50 п.п. к августу",
    }
    assert get_monthly_metrics("ООО Доставка Плюс", "2026-09", db_path=db)["total_spend"] == 33_000.0


def test_history_is_off_by_default_and_skips_mock_data(tmp_path, db, caplog):
    generate_report(mock=True, output=tmp_path, mock_seed=1, today=dt.date(2026, 9, 22))
    with caplog.at_level(logging.INFO, logger="src"):
        result = generate_report(mock=True, output=tmp_path, history_db=db, mock_seed=1, today=dt.date(2026, 9, 22))

    assert (result.report_month, result.compared_to, result.history_saved) == (None, None, False)
    assert not db.exists()
    assert "синтетические данные в неё не записываются" in caplog.text


def test_partial_month_is_saved_with_warning(tmp_path, db, caplog):
    source = tmp_path / "part.csv"
    source.write_text(f"{HEADER}\n2026-09-10,Поиск,100,5,50,1\n2026-09-11,Поиск,100,5,50,1", encoding="utf-8")

    with caplog.at_level(logging.WARNING, logger="src"):
        result = generate_report(input_path=source, output=tmp_path, history_db=db, today=dt.date(2026, 9, 22))

    assert result.report_month == "2026-09" and result.history_saved
    assert "не полный календарный месяц" in caplog.text


def test_broken_history_does_not_break_the_report(tmp_path, caplog):
    broken = tmp_path / "history.db"
    broken.write_bytes(b"not a sqlite database" * 100)
    source = month_csv(tmp_path / "sep.csv", 9, cost=1000.0, clicks=20, conversions=2)

    with caplog.at_level(logging.WARNING, logger="src"):
        result = generate_report(input_path=source, output=tmp_path, history_db=broken, today=dt.date(2026, 10, 1))

    assert result.output_path.exists()
    assert (result.compared_to, result.history_saved) == (None, False)
    assert "История недоступна" in caplog.text and "Не удалось сохранить" in caplog.text


def test_cli_history_flag(tmp_path, db, capsys):
    august = month_csv(tmp_path / "aug.csv", 8, cost=1000.0, clicks=20, conversions=2)
    september = month_csv(tmp_path / "sep.csv", 9, cost=900.0, clicks=20, conversions=3)

    assert cli.main(["-i", str(august), "-o", str(tmp_path), "--history", str(db), "--no-color"]) == 0
    assert cli.main(["-i", str(september), "-o", str(tmp_path), "--history", str(db), "--no-color"]) == 0

    out = capsys.readouterr().out
    assert "история: итоги 2026-08 сохранены · прошлого месяца в истории нет" in out
    assert "история: итоги 2026-09 сохранены · динамика MoM к 2026-08" in out
