import datetime as dt
from pathlib import Path
from typing import get_args

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import PP_ALIGN

from src.config_loader import load_report_config
from src.data_loader import generate_mock_ad_data
from src.rendering.builder import build_presentation
from src.rendering.formatting import NBSP, format_number, format_value
from src.rendering.primitives import CONTENT_WIDTH
from src.rendering.slides import SLIDE_RENDERERS
from src.rendering.slides.notes import paginate_notes
from src.rendering.slides.summary_table import Deviation, deviation, layout_table
from src.rendering.slides.trend_chart import axis_number_format, nice_step
from src.rendering.theme import WHITE, Theme, hex_to_rgb
from src.schemas import ExpertNotes, ReportConfig, Slide, ThemeConfig, aggregate_by_campaign

END = dt.date(2026, 9, 21)
GENERATED = dt.date(2026, 9, 22)


@pytest.fixture(scope="module")
def config() -> ReportConfig:
    return load_report_config()


@pytest.fixture(scope="module")
def records():
    return generate_mock_ad_data(days=30, campaigns_count=4, end_date=END, seed=1)


def build(tmp_path: Path, config: ReportConfig, records, name: str = "report.pptx") -> Presentation:
    path = build_presentation(config, aggregate_by_campaign(records), records, tmp_path / name, generated_at=GENERATED)
    return Presentation(path)


def with_slide(config: ReportConfig, slide_type: str, **changes) -> ReportConfig:
    raw = config.model_dump()
    for slide in raw["slides"]:
        if slide["type"] == slide_type:
            slide.update(changes)
    return ReportConfig.model_validate(raw)


def texts(slide) -> list[str]:
    found = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
    for shape in slide.shapes:
        if shape.has_table:
            found += [cell.text for row in shape.table.rows for cell in row.cells]
    return found


def shape_named(slide, name: str):
    return next(shape for shape in slide.shapes if shape.name == name)


def test_builds_widescreen_deck_with_one_slide_per_active_slide(tmp_path, config, records):
    path = build_presentation(config, aggregate_by_campaign(records), records, tmp_path / "out" / "report.pptx")

    assert path.is_absolute() and path.exists()
    deck = Presentation(path)
    assert len(deck.slides) == len(config.active_slides) == 5
    assert (deck.slide_width, deck.slide_height) == (12_192_000, 6_858_000)  # 13,333″ × 7,5″
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9, rel=1e-3)


def test_disabled_slides_are_skipped(tmp_path, config, records):
    deck = build(tmp_path, with_slide(config, "notes_slide", enabled=False), records)

    assert len(deck.slides) == 4


def test_rejects_non_pptx_path(tmp_path, config, records):
    with pytest.raises(ValueError, match=".pptx"):
        build_presentation(config, [], records, tmp_path / "report.pdf")


def test_every_slide_type_has_a_renderer():
    slide_models = get_args(get_args(Slide)[0])
    config_types = {get_args(model.model_fields["type"].annotation)[0] for model in slide_models}

    assert config_types == set(SLIDE_RENDERERS)


def test_title_slide_shows_client_period_and_generation_date(tmp_path, config, records):
    content = texts(build(tmp_path, config, records).slides[0])

    assert "ООО Доставка Плюс" in content
    assert "Период: 23.08.2026 – 21.09.2026" in content
    assert "Дата формирования: 22.09.2026" in content


def test_kpi_cards_show_formatted_totals(tmp_path, config, records):
    slide = build(tmp_path, config, records).slides[1]
    total_cost = sum(record.cost for record in records)

    spend = shape_named(slide, "KPI total_spend: значение").text_frame.text
    assert spend == f"{format_number(total_cost)}{NBSP}₽"
    assert shape_named(slide, "KPI avg_ctr: значение").text_frame.text.endswith("%")


def test_summary_table_has_totals_row_and_aligned_columns(tmp_path, config, records):
    table = shape_named(build(tmp_path, config, records).slides[2], "Сводная таблица").table
    rows = list(table.rows)

    assert len(rows) == 4 + 2  # кампании + шапка + «ИТОГО»
    assert rows[-1].cells[0].text == "ИТОГО"
    name_cell, cost_cell = rows[1].cells[0], rows[1].cells[1]
    assert name_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT
    assert cost_cell.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    assert sum(column.width for column in table.columns) == CONTENT_WIDTH


def test_summary_table_collapses_tail_into_others_row(tmp_path, config):
    many = generate_mock_ad_data(days=7, campaigns_count=15, end_date=END, seed=2)
    table = shape_named(build(tmp_path, config, many).slides[2], "Сводная таблица").table
    names = [row.cells[0].text for row in table.rows]

    assert len(names) == 10 + 2
    assert names[-2] == "Прочие кампании (6)"


@pytest.mark.parametrize("chart_type, expected", [("line", XL_CHART_TYPE.LINE), ("bar", XL_CHART_TYPE.COLUMN_CLUSTERED)])
def test_trend_chart_uses_native_charts(tmp_path, config, records, chart_type, expected):
    deck = build(tmp_path, with_slide(config, "trend_chart", chart_type=chart_type), records)
    charts = [shape.chart for shape in deck.slides[3].shapes if shape.has_chart]

    assert len(charts) == 2  # Cost и Conversions — у каждой своя ось
    assert all(chart.chart_type == expected for chart in charts)
    assert all(len(list(chart.plots[0].categories)) == 30 for chart in charts)
    assert all("," not in chart.value_axis.tick_labels.number_format for chart in charts)  # без «40,000»


def test_notes_slide_has_placeholder_text(tmp_path, config, records):
    slide = build(tmp_path, config, records).slides[4]

    assert "Вставьте комментарии специалиста" in shape_named(slide, "Комментарий специалиста").text_frame.text


def test_empty_data_renders_without_errors(tmp_path, config):
    deck = build(tmp_path, config, [])

    assert len(deck.slides) == 5
    assert "—" in texts(deck.slides[1])  # CPA и CTR не определены


@pytest.mark.parametrize(
    ("row", "benchmark", "expected"),
    [
        ({"CPA": 1500, "Cost": 1}, 1000, Deviation.WORSE),
        ({"CPA": 700, "Cost": 1}, 1000, Deviation.BETTER),
        ({"CPA": 1100, "Cost": 1}, 1000, None),  # в пределах 20 %
        ({"CPA": None, "Cost": 5000}, 1000, Deviation.WORSE),  # расход без конверсий
        ({"CPA": None, "Cost": 0}, 1000, None),
        ({"CTR": 3.0}, 2.0, Deviation.BETTER),
    ],
)
def test_deviation_respects_metric_direction(row, benchmark, expected):
    field = next(iter(row))

    assert deviation(field, row, benchmark) == expected


def test_cost_deviation_is_neutral():
    assert deviation("Cost", {"Cost": 5000}, 1000) == Deviation.NOTABLE


def test_table_layout_fills_width_and_shrinks_font_for_wide_tables():
    fields = ["CampaignName", "Cost"]
    narrow = layout_table(fields, ["Кампания", "Расход"], [["Поиск", "1 000"]], ["ИТОГО", "1 000"], 10**8)
    wide_fields = ["CampaignId", "CampaignName", *["Impressions"] * 9]
    wide = layout_table(
        wide_fields, ["ID", "Кампания", *["Показы"] * 9],
        [["111991378", "РСЯ_Москва_Похожие_Покупатели", *["1 800 000"] * 9]],
        ["", "ИТОГО", *["3 744 103"] * 9], 10**8,
    )  # fmt: skip

    assert sum(narrow.column_widths) == sum(wide.column_widths) == CONTENT_WIDTH
    assert narrow.font_size > wide.font_size


def test_format_value():
    assert format_value(1234567.891, "currency", "₽") == f"1{NBSP}234{NBSP}568{NBSP}₽"
    assert format_value(2.456, "percent", "₽", decimals=2) == "2,46%"
    assert format_value(None, "currency", "₽") == "—"


def test_hex_to_rgb():
    assert hex_to_rgb("#1a365d") == RGBColor(0x1A, 0x36, 0x5D)
    with pytest.raises(ValueError):
        hex_to_rgb("1A365")


def test_text_color_follows_background_contrast():
    theme = Theme.from_config(ThemeConfig())

    assert theme.text_on(theme.primary) == WHITE
    assert theme.text_on(hex_to_rgb("F8F9FA")) == theme.text


@pytest.mark.parametrize(("peak", "step"), [(38_500, 10_000), (2.4, 1), (0.9, 0.5), (0, 1)])
def test_nice_step(peak, step):
    assert nice_step(peak) == pytest.approx(step)


def test_axis_format_is_locale_independent_for_whole_steps():
    assert axis_number_format(10_000, percent=False) == "[>=1000000]0 000 000;[>=1000]0 000;0"
    assert axis_number_format(1, percent=True) == '0"%"'
    assert axis_number_format(0.5, percent=True) == '0.0"%"'


def test_notes_pagination_keeps_heading_with_following_block():
    items = "\n".join(f"- Пункт {index}: подробный комментарий специалиста по кампании и ставкам" for index in range(30))
    notes = ExpertNotes.parse(f"# Разбор\n{items}\n# План\n1. Действие")

    pages = paginate_notes(notes)

    assert len(pages) > 1
    assert sum(len(page.blocks) for page in pages) == len(notes.blocks)
    assert all(page.blocks[-1].kind != "heading" for page in pages)


def test_short_notes_fit_one_slide_with_bold_runs(tmp_path, config, records):
    notes = ExpertNotes.parse("Итог: **CPA −12 %** за месяц")
    path = build_presentation(config, aggregate_by_campaign(records), records, tmp_path / "r.pptx", notes=notes)
    shape = shape_named(Presentation(path).slides[4], "Комментарий специалиста")

    runs = shape.text_frame.paragraphs[0].runs
    assert [(run.text, bool(run.font.bold)) for run in runs] == [("Итог: ", False), ("CPA −12 %", True), (" за месяц", False)]