"""Сборка презентации PowerPoint (16:9) по ReportConfig.

Проверочный запуск из корня проекта — моковые данные за период из конфига:
    python -m src.rendering.builder
    python src/rendering/builder.py
"""

import datetime as dt
import logging
import sys
from collections.abc import Sequence
from pathlib import Path

if not __package__:  # запуск файлом: делаем пакет src видимым
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from pptx import Presentation  # noqa: E402
from pptx.presentation import Presentation as PresentationType  # noqa: E402
from pptx.slide import SlideLayout  # noqa: E402

from src.config_loader import PROJECT_ROOT, load_report_config  # noqa: E402
from src.console import configure_logging, use_utf8_console  # noqa: E402
from src.data_loader import generate_mock_ad_data  # noqa: E402
from src.rendering.context import MonthComparison, RenderContext  # noqa: E402
from src.rendering.primitives import SLIDE_HEIGHT, SLIDE_WIDTH, add_footer, fill_background  # noqa: E402
from src.rendering.slides import plan_slides, render_slide  # noqa: E402
from src.schemas import (  # noqa: E402
    CampaignSummary,
    DailyAdRecord,
    ExpertNotes,
    ReportConfig,
    aggregate_by_campaign,
)

OUTPUT_DIR = PROJECT_ROOT / "output"
TEST_REPORT_PATH = OUTPUT_DIR / "report_test.pptx"

logger = logging.getLogger(__name__)


def build_presentation(
    config: ReportConfig,
    summary: Sequence[CampaignSummary],
    daily_records: Sequence[DailyAdRecord],
    output_path: Path,
    *,
    generated_at: dt.date | None = None,
    notes: ExpertNotes | None = None,
    comparison: MonthComparison | None = None,
) -> Path:
    """Собирает презентацию из активных слайдов конфига и сохраняет её.

    Args:
        config: провалидированный конфиг отчёта.
        summary: сводки по кампаниям — таблица, KPI и строка «ИТОГО».
        daily_records: суточные записи — динамика по дням и период отчёта.
        output_path: путь к .pptx; недостающие папки создаются.
        generated_at: дата формирования на титульном слайде, по умолчанию сегодня.
        notes: выводы специалиста для слайда notes_slide; без них там будет заготовка.
        comparison: итоги прошлого месяца из истории — карточки KPI покажут динамику MoM.

    Returns:
        Абсолютный путь к сохранённому файлу.
    """
    output_path = Path(output_path)
    if output_path.suffix.lower() != ".pptx":
        raise ValueError(f"Ожидается путь к файлу .pptx, получено: {output_path}")

    ctx = RenderContext.from_data(config, summary, daily_records, generated_at or dt.date.today(), notes, comparison)
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    layout = _blank_layout(presentation)

    planned = plan_slides(config, ctx)  # длинные выводы могут занять несколько слайдов
    for number, (slide_config, slide_ctx) in enumerate(planned, start=1):
        slide = presentation.slides.add_slide(layout)
        fill_background(slide, ctx.theme.background)
        render_slide(slide, slide_config, slide_ctx)
        logger.debug("Слайд %d/%d: %s «%s»", number, len(planned), slide_config.type, slide_config.title)
        if slide_config.type != "title_slide":
            add_footer(slide, ctx.theme, f"{ctx.client_name} · {ctx.period_label}", f"{number} / {len(planned)}")

    properties = presentation.core_properties
    properties.title = f"{ctx.client_name}: отчёт за {ctx.period_label}"
    properties.author = "Генератор отчётов"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    return output_path.resolve()


def _blank_layout(presentation: PresentationType) -> SlideLayout:
    """Пустой макет шаблона: всё содержимое слайда рисуют рендереры."""
    for layout in presentation.slide_layouts:
        if layout.name == "Blank":
            return layout
    return presentation.slide_layouts[6]


def main() -> None:
    """Проверочный прогон: конфиг → моковые данные за его период → output/report_test.pptx."""
    use_utf8_console()
    configure_logging()
    config = load_report_config()
    date_from, date_to = config.report_metadata.resolve_period()
    records = generate_mock_ad_data(days=(date_to - date_from).days + 1, campaigns_count=4, end_date=date_to)
    summaries = aggregate_by_campaign(records)
    try:
        path = build_presentation(config, summaries, records, TEST_REPORT_PATH)
    except PermissionError:
        raise SystemExit(f"Не удалось записать {TEST_REPORT_PATH}: файл открыт в PowerPoint?") from None

    logger.info("Презентация сохранена: %s", path)
    logger.info(
        "Слайдов: %d, кампаний: %d, период: %s – %s",
        len(config.active_slides), len(summaries), f"{date_from:%d.%m.%Y}", f"{date_to:%d.%m.%Y}",
    )  # fmt: skip


if __name__ == "__main__":
    main()
