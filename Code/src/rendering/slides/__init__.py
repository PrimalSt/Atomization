"""Фабрика слайдов: тип слайда из конфига → функция-рендерер.

Новый тип слайда = модель в ``src.schemas`` + модуль здесь + строка в ``SLIDE_RENDERERS``.
Если содержимое может не поместиться на один слайд, тип регистрирует ещё и
«расширитель» в ``SLIDE_EXPANDERS``: он превращает слайд конфига в несколько.
"""

from collections.abc import Callable
from typing import Any

from pptx.slide import Slide

from src.rendering.context import RenderContext
from src.rendering.slides.kpi_cards import render_kpi_cards
from src.rendering.slides.notes import expand_notes_slide, render_notes_slide
from src.rendering.slides.summary_table import render_summary_table
from src.rendering.slides.title import render_title_slide
from src.rendering.slides.trend_chart import render_trend_chart
from src.schemas import ReportConfig
from src.schemas import Slide as SlideConfig

SlideRenderer = Callable[[Slide, Any, RenderContext], None]
SlideExpander = Callable[[Any, RenderContext], list[tuple[Any, RenderContext]]]

SLIDE_RENDERERS: dict[str, SlideRenderer] = {
    "title_slide": render_title_slide,
    "kpi_cards": render_kpi_cards,
    "summary_table": render_summary_table,
    "trend_chart": render_trend_chart,
    "notes_slide": render_notes_slide,
}

SLIDE_EXPANDERS: dict[str, SlideExpander] = {
    "notes_slide": expand_notes_slide,
}


def plan_slides(config: ReportConfig, ctx: RenderContext) -> list[tuple[SlideConfig, RenderContext]]:
    """Итоговый список слайдов презентации: активные слайды конфига с учётом продолжений."""
    planned: list[tuple[SlideConfig, RenderContext]] = []
    for slide_config in config.active_slides:
        expander = SLIDE_EXPANDERS.get(slide_config.type)
        planned.extend(expander(slide_config, ctx) if expander else [(slide_config, ctx)])
    return planned


def render_slide(slide: Slide, config: SlideConfig, ctx: RenderContext) -> None:
    try:
        renderer = SLIDE_RENDERERS[config.type]
    except KeyError:
        raise ValueError(f"Нет рендерера для слайда типа «{config.type}»") from None
    renderer(slide, config, ctx)
