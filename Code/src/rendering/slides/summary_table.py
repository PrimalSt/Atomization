"""Таблица по кампаниям со строкой «ИТОГО» и подсветкой отклонений выбранной метрики."""

import math
from collections.abc import Sequence
from dataclasses import dataclass
from enum import Enum
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.table import _Cell
from pptx.util import Inches, Pt

from src.rendering.context import RenderContext
from src.rendering.formatting import FIELD_SPECS, RATIO_FIELDS, format_field
from src.rendering.primitives import (
    CONTENT_HEIGHT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    MARGIN,
    add_slide_title,
    add_text,
    set_font,
)
from src.schemas import CampaignSummary, PerformanceMetrics, SummaryTableSlide, calculate_totals

MAX_CAMPAIGN_ROWS = 10  # кампании сверх лимита сворачиваются в строку «Прочие кампании»
HIGHLIGHT_THRESHOLD = 0.2  # подсвечиваем отклонение от ориентира от 20 %

_FONT_SIZES = (14, 12, 11, 10)
_CELL_MARGIN_X = Inches(0.1)
_CELL_MARGIN_Y = Inches(0.03)
_CHAR_WIDTH_EM = {False: 0.58, True: 0.66}  # средняя ширина символа Calibri с запасом: обычный / жирный
_MIN_COLUMN_WIDTH = Inches(1.0)  # короткие колонки («Лиды», «CR») не должны быть впритык
_MIN_STRETCH_WIDTH = Inches(1.6)  # уже этого название кампании нечитаемо
_LEGEND_GAP = Inches(0.15)
_LEGEND_HEIGHT = Inches(0.3)

# +1 — чем больше, тем лучше; −1 — чем меньше, тем лучше; 0 — без оценки (расход).
_DIRECTION: dict[str, int] = {
    "Impressions": 1,
    "Clicks": 1,
    "Conversions": 1,
    "CTR": 1,
    "CR": 1,
    "CPC": -1,
    "CPA": -1,
    "Cost": 0,
}


class Deviation(Enum):
    BETTER = "better"
    WORSE = "worse"
    NOTABLE = "notable"  # заметное отклонение метрики без «хорошего» направления


def deviation(field: str, row: dict[str, Any], benchmark: float | None) -> Deviation | None:
    """Оценивает значение поля в строке относительно ориентира.

    Особый случай: CPA не определён, но расход есть — кампания тратит бюджет
    без конверсий, это худший возможный CPA.
    """
    value = row.get(field)
    if value is None:
        return Deviation.WORSE if field == "CPA" and row.get("Cost", 0) > 0 else None
    if not benchmark:
        return None
    change = value / benchmark - 1
    if abs(change) < HIGHLIGHT_THRESHOLD:
        return None
    direction = _DIRECTION[field]
    if direction == 0:
        return Deviation.NOTABLE
    return Deviation.BETTER if (change > 0) == (direction > 0) else Deviation.WORSE


def benchmark_for(field: str, totals: PerformanceMetrics, campaigns_count: int) -> float | None:
    """Ориентир: для производных метрик — значение по аккаунту, для счётчиков — среднее на кампанию."""
    value = totals.model_dump(by_alias=True)[field]
    if value is None or field in RATIO_FIELDS:
        return value
    return value / campaigns_count if campaigns_count else None


def campaign_rows(summaries: Sequence[CampaignSummary]) -> list[dict[str, Any]]:
    """Строки кампаний по убыванию расхода; хвост сверх лимита — одной строкой «Прочие»."""
    ordered = sorted(summaries, key=lambda summary: (-summary.cost, summary.campaign_id))
    if len(ordered) <= MAX_CAMPAIGN_ROWS:
        return [summary.model_dump(by_alias=True) for summary in ordered]
    shown, rest = ordered[: MAX_CAMPAIGN_ROWS - 1], ordered[MAX_CAMPAIGN_ROWS - 1 :]
    others = calculate_totals(rest).model_dump(by_alias=True) | {
        "CampaignName": f"Прочие кампании ({len(rest)})",
        "CampaignId": "",
        "DaysCount": "",
    }
    return [*(summary.model_dump(by_alias=True) for summary in shown), others]


@dataclass(frozen=True, slots=True)
class TableLayout:
    font_size: float
    column_widths: list[int]
    row_height: int


def layout_table(
    fields: Sequence[str],
    header: Sequence[str],
    body: Sequence[Sequence[str]],
    total: Sequence[str],
    height_limit: int,
) -> TableLayout:
    """Подбирает кегль и ширины колонок под содержимое.

    Числовые колонки получают ширину по самому длинному значению — число не
    разрывается переносом. Свободное место сначала доводит узкие колонки
    до минимальной ширины, остаток отдаётся названию кампании. Берётся самый
    крупный кегль, при котором таблица помещается по ширине и высоте; если не
    помещается и на минимальном, названия переносятся, а строки становятся выше.
    """
    rows_count = len(body) + 2
    for size in _FONT_SIZES:
        natural = _natural_widths(header, body, total, size)
        if sum(natural) <= CONTENT_WIDTH and rows_count * _row_height(size, 1) <= height_limit:
            break

    stretch = [index for index, field in enumerate(fields) if field == "CampaignName"]
    stretch = stretch or list(range(len(fields)))
    widths = list(natural)
    spare = CONTENT_WIDTH - sum(natural)
    if spare >= 0:
        for index in range(len(widths)):
            if index not in stretch:
                extra = min(spare, max(0, _MIN_COLUMN_WIDTH - widths[index]))
                widths[index] += extra
                spare -= extra
        for index in stretch:
            widths[index] += spare // len(stretch)
    else:
        for index in stretch:
            widths[index] = max(_MIN_STRETCH_WIDTH, widths[index] + spare // len(stretch))
        if sum(widths) > CONTENT_WIDTH:  # не помещается и так — сжимаем все колонки пропорционально
            widths = [width * CONTENT_WIDTH // sum(widths) for width in widths]
    widths[-1] += CONTENT_WIDTH - sum(widths)

    lines = max(math.ceil(natural[index] / widths[index]) for index in range(len(fields)))
    return TableLayout(font_size=size, column_widths=widths, row_height=_row_height(size, lines))


def _text_width(text: str, size: float, bold: bool) -> int:
    """Оценка ширины строки: python-pptx не измеряет текст, считаем по средней ширине символа."""
    return int(Pt(size) * _CHAR_WIDTH_EM[bold] * len(text))


def _natural_widths(
    header: Sequence[str], body: Sequence[Sequence[str]], total: Sequence[str], size: float
) -> list[int]:
    widths = []
    for column, title in enumerate(header):
        content = max((_text_width(row[column], size, bold=False) for row in body), default=0)
        widest = max(_text_width(title, size, bold=True), _text_width(total[column], size, bold=True), content)
        widths.append(widest + 2 * _CELL_MARGIN_X)
    return widths


def _row_height(size: float, lines: int) -> int:
    return int(Pt(size) * 1.25 * lines + Pt(size) * 1.1 + 2 * _CELL_MARGIN_Y)


def render_summary_table(slide: Slide, config: SummaryTableSlide, ctx: RenderContext) -> None:
    theme = ctx.theme
    add_slide_title(slide, config.title, theme)

    fields = [column.field for column in config.columns]
    rows = campaign_rows(ctx.summaries)
    total_row = ctx.totals.model_dump(by_alias=True) | {"CampaignName": "ИТОГО", "CampaignId": "", "DaysCount": ""}
    header = [column.header for column in config.columns]
    body = [[format_field(field, row[field]) for field in fields] for row in rows]
    total = [format_field(field, total_row[field]) for field in fields]

    highlight = config.highlight_metric
    height_limit = CONTENT_HEIGHT - (_LEGEND_GAP + _LEGEND_HEIGHT if highlight else 0)
    layout = layout_table(fields, header, body, total, height_limit)
    rows_count = len(rows) + 2  # + шапка и «ИТОГО»
    table_height = layout.row_height * rows_count

    frame = slide.shapes.add_table(rows_count, len(fields), MARGIN, CONTENT_TOP, CONTENT_WIDTH, table_height)
    frame.name = "Сводная таблица"
    table = frame.table
    table.horz_banding = False
    for table_row in table.rows:
        table_row.height = layout.row_height
    for table_column, width in zip(table.columns, layout.column_widths):
        table_column.width = width

    size = layout.font_size
    numeric = [FIELD_SPECS[field].numeric for field in fields]
    header_color = theme.text_on(theme.primary)
    for column, text in enumerate(header):
        _fill_cell(table.cell(0, column), text, ctx, fill=theme.primary, color=header_color,
                   size=size, bold=True, numeric=numeric[column])  # fmt: skip

    benchmark = benchmark_for(highlight, ctx.totals, len(ctx.summaries)) if highlight else None
    for row_index, (row, texts) in enumerate(zip(rows, body), start=1):
        zebra = theme.surface if row_index % 2 else theme.primary_soft
        is_campaign = row["CampaignId"] != ""  # «Прочие кампании» — агрегат, его не оцениваем
        for column, (field, text) in enumerate(zip(fields, texts)):
            status = deviation(field, row, benchmark) if field == highlight and is_campaign else None
            fill, color = _status_colors(status, zebra, ctx)
            _fill_cell(table.cell(row_index, column), text, ctx, fill=fill, color=color, size=size,
                       bold=status is not None, numeric=numeric[column])  # fmt: skip

    for column, text in enumerate(total):
        _fill_cell(table.cell(rows_count - 1, column), text, ctx, fill=theme.accent_soft, color=theme.text,
                   size=size, bold=True, numeric=numeric[column])  # fmt: skip

    if highlight:
        add_text(
            slide, MARGIN, CONTENT_TOP + table_height + _LEGEND_GAP, CONTENT_WIDTH, _LEGEND_HEIGHT,
            _legend_text(highlight),
            font=theme.body_font, size=11, color=theme.muted, name="Легенда подсветки",
        )  # fmt: skip


def _fill_cell(
    cell: _Cell,
    text: str,
    ctx: RenderContext,
    *,
    fill: RGBColor,
    color: RGBColor,
    size: float,
    bold: bool,
    numeric: bool,
) -> None:
    """Числа — по правому краю, текст — по левому."""
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = _CELL_MARGIN_X
    cell.margin_top = cell.margin_bottom = _CELL_MARGIN_Y
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT if numeric else PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    set_font(run.font, name=ctx.theme.body_font, size=size, color=color, bold=bold)


def _status_colors(status: Deviation | None, zebra: RGBColor, ctx: RenderContext) -> tuple[RGBColor, RGBColor]:
    theme = ctx.theme
    match status:
        case Deviation.BETTER:
            return theme.good_fill, theme.good
        case Deviation.WORSE:
            return theme.bad_fill, theme.bad
        case Deviation.NOTABLE:
            return theme.accent_soft, theme.primary
    return zebra, theme.text


def _legend_text(field: str) -> str:
    threshold = f"{HIGHLIGHT_THRESHOLD:.0%}".replace("%", " %")
    if _DIRECTION[field] == 0:
        return f"Подсветка {field}: отклонение от среднего на кампанию на {threshold} и больше."
    benchmark = "значения по аккаунту" if field in RATIO_FIELDS else "среднего на кампанию"
    legend = f"Подсветка {field}: зелёный — лучше {benchmark} на {threshold} и больше, красный — хуже."
    if field == "CPA":
        legend += " «—» на красном — расход без конверсий."
    return legend
