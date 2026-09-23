"""Слайд выводов специалиста: заметки эксперта или текст-заготовка."""

import dataclasses
import logging
import math
import re

from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.slide import Slide
from pptx.text.text import _Paragraph
from pptx.util import Inches, Length, Pt

from src.rendering.context import RenderContext
from src.rendering.primitives import (
    CARD_PADDING,
    CONTENT_HEIGHT,
    CONTENT_TOP,
    CONTENT_WIDTH,
    MARGIN,
    add_card,
    add_slide_title,
    add_text,
    set_font,
)
from src.rendering.theme import Theme
from src.schemas import ExpertNotes, NoteBlock, NotesSlide

logger = logging.getLogger(__name__)

DEFAULT_PLACEHOLDER = "Добавьте выводы и план оптимизации перед отправкой клиенту."
NOTES_SHAPE_NAME = "Комментарий специалиста"  # по имени поле легко найти в области выделения

# (кегль, колонок) — от просторного к плотному; берётся первый вариант, в который помещается текст.
# Не помещается и в последний — выводы продолжаются на следующем слайде.
_LAYOUTS = ((18, 1), (16, 1), (14, 1), (14, 2))
_OVERSIZE_LAYOUT = (12, 2)  # для одного блока, который не влезает на слайд целиком
_COLUMN_GAP = Inches(0.5)
_PADDING = CARD_PADDING + Inches(0.1)
_BOX = (MARGIN + _PADDING, CONTENT_TOP + _PADDING, CONTENT_WIDTH - 2 * _PADDING, CONTENT_HEIGHT - 2 * _PADDING)
_CHAR_WIDTH_EM = 0.48  # средняя ширина символа Calibri (замер по рендеру PowerPoint: ≈ 0,46)
_WRAP_LOSS = 0.93  # перенос по словам оставляет конец строки пустым
_COLUMN_BALANCE = 1.1  # колонка рвётся на границе строки, а не ровно по высоте
_LINE_HEIGHT = 1.35  # межстрочный 1,1 × собственный интервал шрифта
_HEADING_SCALE = 1.15
_LIST_INDENT_EM = 1.4
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_LIST_KINDS = {"bullet", "numbered"}
_AFTER_BULLET = ("a:tabLst", "a:defRPr", "a:extLst")  # элементы, которые по схеме идут после маркера


def expand_notes_slide(config: NotesSlide, ctx: RenderContext) -> list[tuple[NotesSlide, RenderContext]]:
    """Разбивает длинные выводы на несколько слайдов «… (продолжение)» с единой раскладкой."""
    if ctx.notes is None:
        return [(config, ctx)]
    pages = paginate_notes(ctx.notes)
    if len(pages) == 1:
        return [(config, ctx)]
    logger.info("Выводы специалиста не помещаются на один слайд — продолжены ещё на %d", len(pages) - 1)
    continued = config.model_copy(update={"title": f"{config.title} (продолжение)"})
    return [
        (config if index == 0 else continued, dataclasses.replace(ctx, notes=page, notes_layout=_LAYOUTS[-1]))
        for index, page in enumerate(pages)
    ]


def paginate_notes(notes: ExpertNotes) -> list[ExpertNotes]:
    """Раскладывает блоки по слайдам так, чтобы каждый помещался в самой плотной раскладке.

    Заголовок не остаётся последним на слайде — он переносится вместе со следующим блоком.
    """
    _, _, width, height = _BOX
    pages: list[list[NoteBlock]] = [[]]
    for block in notes.blocks:
        current = pages[-1]
        if not current or _fits([*current, block], *_LAYOUTS[-1], width, height):
            current.append(block)
            continue
        carried = [current.pop()] if current[-1].kind == "heading" else []
        if current:
            pages.append([*carried, block])
        else:
            current.extend([*carried, block])
    return [ExpertNotes(blocks=page) for page in pages]


def render_notes_slide(slide: Slide, config: NotesSlide, ctx: RenderContext) -> None:
    theme = ctx.theme
    add_slide_title(slide, config.title, theme)
    add_card(slide, MARGIN, CONTENT_TOP, CONTENT_WIDTH, CONTENT_HEIGHT, theme)

    if ctx.notes is None:
        add_text(
            slide, *_BOX, config.placeholder_text or DEFAULT_PLACEHOLDER,
            font=theme.body_font, size=18, color=theme.muted, italic=True, name=NOTES_SHAPE_NAME,
        )  # fmt: skip
        # Напоминание видно только докладчику и не попадает на экран клиенту.
        slide.notes_slide.notes_text_frame.text = (
            f"Замените текст в поле «{NOTES_SHAPE_NAME}» выводами специалиста перед отправкой отчёта."
        )
        return
    _render_notes(slide, ctx.notes, theme, ctx.notes_layout, config.title)


def _render_notes(
    slide: Slide, notes: ExpertNotes, theme: Theme, layout: tuple[int, int] | None, title: str
) -> None:
    left, top, width, height = _BOX
    if layout is None or not _fits(notes.blocks, *layout, width, height):
        layout = _choose_layout(notes, width, height, title)
    size, columns = layout

    textbox = slide.shapes.add_textbox(left, top, width, height)
    textbox.name = NOTES_SHAPE_NAME
    frame = textbox.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    if columns > 1:
        body = frame._bodyPr  # noqa: SLF001 — колонок текста нет в публичном API python-pptx
        body.set("numCol", str(columns))
        body.set("spcCol", str(_COLUMN_GAP))

    for index, block in enumerate(notes.blocks):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _format_block(paragraph, block, size, theme, first=index == 0)


def _choose_layout(notes: ExpertNotes, width: int, height: int, title: str) -> tuple[int, int]:
    """Самый крупный кегль (и одна колонка, если хватает) без выхода текста за карточку."""
    for size, columns in _LAYOUTS:
        if _fits(notes.blocks, size, columns, width, height):
            return size, columns
    size, columns = _OVERSIZE_LAYOUT
    if not _fits(notes.blocks, size, columns, width, height):
        logger.warning(
            "Абзац выводов на слайде «%s» не помещается даже кеглем %d в %d колонки — сократите его",
            title, size, columns,
        )  # fmt: skip
    return size, columns


def _fits(blocks: list[NoteBlock], size: int, columns: int, width: int, height: int) -> bool:
    column_width = (width - _COLUMN_GAP * (columns - 1)) // columns
    needed = _estimate_height(blocks, size, column_width) / columns * (_COLUMN_BALANCE if columns > 1 else 1)
    return needed <= Length(height).pt


def _estimate_height(blocks: list[NoteBlock], size: int, width: int) -> float:
    """Высота текста в пунктах по средней ширине символа (python-pptx текст не измеряет)."""
    total = 0.0
    for index, block in enumerate(blocks):
        block_size = _block_size(block, size)
        indent = Pt(block_size) * _LIST_INDENT_EM if block.kind in _LIST_KINDS else 0
        chars_per_line = max(1, int((width - indent) / (Pt(block_size) * _CHAR_WIDTH_EM) * _WRAP_LOSS))
        lines = math.ceil(len(_BOLD.sub(r"\1", block.text)) / chars_per_line)
        before, after = _spacing(block, size, first=index == 0)
        total += lines * block_size * _LINE_HEIGHT + before + after
    return total


def _block_size(block: NoteBlock, size: int) -> float:
    return size * _HEADING_SCALE if block.kind == "heading" else size


def _spacing(block: NoteBlock, size: int, *, first: bool) -> tuple[float, float]:
    """Интервалы до и после абзаца, пт: заголовок отбивается сверху, пункты списка — плотнее."""
    if block.kind == "heading":
        return (0 if first else size * 0.7), size * 0.3
    return 0, size * (0.3 if block.kind in _LIST_KINDS else 0.6)


def _format_block(paragraph: _Paragraph, block: NoteBlock, size: int, theme: Theme, *, first: bool) -> None:
    heading = block.kind == "heading"
    block_size = _block_size(block, size)
    before, after = _spacing(block, size, first=first)
    paragraph.line_spacing = 1.1
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    for text, bold in _runs(block.text):
        run = paragraph.add_run()
        run.text = text
        set_font(
            run.font, name=theme.heading_font if heading else theme.body_font, size=block_size,
            color=theme.primary if heading else theme.text, bold=heading or bold,
        )  # fmt: skip
    if block.kind in _LIST_KINDS:
        _set_list_marker(paragraph, numbered=block.kind == "numbered", size=block_size, color=theme.primary)


def _runs(text: str) -> list[tuple[str, bool]]:
    """«Итог: **CPA −12 %**» → [(«Итог: », обычный), («CPA −12 %», полужирный)]."""
    parts = _BOLD.split(text)
    return [(part, index % 2 == 1) for index, part in enumerate(parts) if part]


def _set_list_marker(paragraph: _Paragraph, *, numbered: bool, size: float, color: RGBColor) -> None:
    """Настоящий маркер списка PowerPoint с висячим отступом — не символ «•» в тексте.

    В python-pptx нет API для маркеров, поэтому элементы a:buClr, a:buFont и
    a:buChar / a:buAutoNum добавляются в свойства абзаца в порядке, заданном схемой.
    """
    indent = int(Pt(size) * _LIST_INDENT_EM)
    properties = paragraph._p.get_or_add_pPr()  # noqa: SLF001 — см. докстринг
    properties.set("marL", str(indent))
    properties.set("indent", str(-indent))
    markers = [f'<a:buClr {nsdecls("a")}><a:srgbClr val="{color}"/></a:buClr>']
    if numbered:
        markers.append(f'<a:buAutoNum {nsdecls("a")} type="arabicPeriod"/>')
    else:
        markers.append(f'<a:buFont {nsdecls("a")} typeface="Arial"/>')
        markers.append(f'<a:buChar {nsdecls("a")} char="•"/>')
    for marker in markers:
        properties.insert_element_before(parse_xml(marker), *_AFTER_BULLET)
