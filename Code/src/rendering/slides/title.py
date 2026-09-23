"""Титульный слайд: название отчёта, клиент, период, дата формирования."""

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Inches, Pt

from src.rendering.context import RenderContext
from src.rendering.primitives import SLIDE_WIDTH, add_text, fill_background, set_font
from src.rendering.theme import blend
from src.schemas import TitleSlide

_LEFT = Inches(0.9)
_WIDTH = SLIDE_WIDTH - 2 * _LEFT
_CHIP_FONT_SIZE = 16


def render_title_slide(slide: Slide, config: TitleSlide, ctx: RenderContext) -> None:
    theme = ctx.theme
    on_primary = theme.text_on(theme.primary)
    secondary = blend(on_primary, theme.primary, 0.3)  # приглушённый текст на тёмном фоне

    fill_background(slide, theme.primary)
    add_text(
        slide, _LEFT, Inches(1.2), _WIDTH, Inches(0.5), ctx.client_name,
        font=theme.body_font, size=20, color=theme.accent_on_primary, bold=True, name="Клиент",
    )  # fmt: skip
    add_text(
        slide, _LEFT, Inches(1.8), _WIDTH, Inches(1.7), config.title,
        font=theme.heading_font, size=48, color=on_primary, bold=True,
        anchor=MSO_ANCHOR.BOTTOM, name="Название отчёта",
    )  # fmt: skip
    if config.subtitle:
        add_text(
            slide, _LEFT, Inches(3.65), _WIDTH, Inches(0.6), config.subtitle,
            font=theme.body_font, size=24, color=secondary, name="Подзаголовок",
        )  # fmt: skip

    _add_period_chip(slide, f"Период: {ctx.period_label}", ctx)
    add_text(
        slide, _LEFT, Inches(6.3), _WIDTH, Inches(0.4),
        f"Дата формирования: {ctx.generated_at:%d.%m.%Y}",
        font=theme.body_font, size=14, color=secondary, name="Дата формирования",
    )  # fmt: skip


def _add_period_chip(slide: Slide, text: str, ctx: RenderContext) -> None:
    """Плашка с периодом: ширина подстраивается под длину текста."""
    theme = ctx.theme
    width = Inches(0.6) + int(Pt(_CHIP_FONT_SIZE) * 0.6 * len(text))
    chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _LEFT, Inches(4.7), width, Inches(0.55))
    chip.name = "Период"
    chip.adjustments[0] = 0.5  # полностью скруглённые торцы
    chip.fill.solid()
    chip.fill.fore_color.rgb = theme.accent
    chip.line.fill.background()
    chip.shadow.inherit = False

    frame = chip.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.word_wrap = False
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    set_font(
        run.font, name=theme.body_font, size=_CHIP_FONT_SIZE, color=theme.text_on(theme.accent), bold=True
    )
