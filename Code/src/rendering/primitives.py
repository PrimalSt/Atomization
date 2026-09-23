"""Геометрия слайда 16:9 и базовые элементы: фон, текст, карточки, заголовок, подвал."""

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.shapes.autoshape import Shape
from pptx.slide import Slide
from pptx.text.text import Font
from pptx.util import Emu, Inches, Pt

from src.rendering.theme import Theme

SLIDE_WIDTH = Emu(12_192_000)  # 13,333″ — стандартный широкоэкранный слайд PowerPoint
SLIDE_HEIGHT = Emu(6_858_000)  # 7,5″

MARGIN = Inches(0.6)
GAP = Inches(0.3)  # единый отступ между блоками
CONTENT_TOP = Inches(1.5)
CONTENT_BOTTOM = Inches(6.7)
CONTENT_WIDTH = SLIDE_WIDTH - 2 * MARGIN
CONTENT_HEIGHT = CONTENT_BOTTOM - CONTENT_TOP
CARD_PADDING = Inches(0.3)

_TITLE_TOP = Inches(0.45)
_TITLE_HEIGHT = Inches(0.8)
_FOOTER_TOP = Inches(6.95)
_FOOTER_HEIGHT = Inches(0.3)
_CARD_RADIUS = Inches(0.12)


def fill_background(slide: Slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def set_font(
    font: Font,
    *,
    name: str,
    size: float,
    color: RGBColor,
    bold: bool = False,
    italic: bool = False,
) -> None:
    font.name = name
    font.size = Pt(size)
    font.color.rgb = color
    font.bold = bold
    font.italic = italic


def add_text(
    slide: Slide,
    left: int,
    top: int,
    width: int,
    height: int,
    text: str,
    *,
    font: str,
    size: float,
    color: RGBColor,
    bold: bool = False,
    italic: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
    name: str | None = None,
) -> Shape:
    """Текстовое поле без внутренних отступов: край текста совпадает с сеткой слайда.

    Каждая строка ``text`` — отдельный абзац с явным форматированием run'а.
    """
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        set_font(run.font, name=font, size=size, color=color, bold=bold, italic=italic)
    if name:
        box.name = name
    return box


def add_card(
    slide: Slide,
    left: int,
    top: int,
    width: int,
    height: int,
    theme: Theme,
    *,
    fill: RGBColor | None = None,
) -> Shape:
    """Карточка со скруглёнными углами и тонкой рамкой — общий мотив контентных слайдов."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = _CARD_RADIUS / min(width, height)  # радиус не зависит от размера
    card.fill.solid()
    card.fill.fore_color.rgb = fill or theme.surface
    card.line.color.rgb = theme.border
    card.line.width = Pt(1)
    card.shadow.inherit = False  # без тени из темы Office
    return card


def add_slide_title(slide: Slide, title: str, theme: Theme) -> None:
    add_text(
        slide, MARGIN, _TITLE_TOP, CONTENT_WIDTH, _TITLE_HEIGHT, title,
        font=theme.heading_font, size=32, color=theme.primary, bold=True,
        anchor=MSO_ANCHOR.MIDDLE, name="Заголовок",
    )  # fmt: skip


def add_footer(slide: Slide, theme: Theme, caption: str, page: str) -> None:
    """Подвал: клиент и период слева, номер слайда справа."""
    page_width = Inches(1.2)
    style = {"font": theme.body_font, "size": 10, "color": theme.muted}
    add_text(
        slide, MARGIN, _FOOTER_TOP, CONTENT_WIDTH - page_width - GAP, _FOOTER_HEIGHT, caption,
        **style, name="Подвал",
    )  # fmt: skip
    add_text(
        slide, MARGIN + CONTENT_WIDTH - page_width, _FOOTER_TOP, page_width, _FOOTER_HEIGHT, page,
        **style, align=PP_ALIGN.RIGHT, name="Номер слайда",
    )  # fmt: skip
